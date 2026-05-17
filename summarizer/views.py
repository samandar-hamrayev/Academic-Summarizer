import io
import json
import logging

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.utils.translation import gettext_lazy as _
from django.views.decorators.http import require_POST
from django.views.generic import ListView

from papers.models import Paper

from .models import ChatMessage, ShareLink, Summary

logger = logging.getLogger(__name__)


# ── AI Tag Preview ────────────────────────────────────────────────────────────

@login_required
@require_POST
def preview_tags(request):
    """
    AJAX endpoint: extract AI-suggested tags from an uploaded PDF without saving it.
    Called by the upload form before submission so the user can confirm/edit tags.
    """
    if 'file' not in request.FILES:
        return JsonResponse({'error': 'No file provided'}, status=400)

    pdf_file = request.FILES['file']

    if not pdf_file.name.lower().endswith('.pdf'):
        return JsonResponse({'error': 'Only PDF files are accepted'}, status=400)

    if pdf_file.size > 50 * 1024 * 1024:
        return JsonResponse({'error': 'File too large (max 50 MB)'}, status=400)

    try:
        from papers.services import extract_text_from_pdf
        text = extract_text_from_pdf(pdf_file)
    except Exception as exc:
        logger.error('Text extraction for tag preview failed: %s', exc)
        return JsonResponse({'suggested_tags': [], 'warning': 'Could not extract text from PDF'})

    if not text or len(text.strip()) < 100:
        return JsonResponse({'suggested_tags': [], 'warning': 'Not enough readable text in PDF'})

    from .services import extract_tags_from_text
    suggested_tags = extract_tags_from_text(text, max_tags=5)
    return JsonResponse({'suggested_tags': suggested_tags})


# ── Summary Views ─────────────────────────────────────────────────────────────

class SummaryHistoryView(LoginRequiredMixin, ListView):
    model = Summary
    template_name = 'summarizer/history.html'
    context_object_name = 'summaries'
    paginate_by = 15

    def get_queryset(self):
        return (
            Summary.objects
            .filter(paper__uploaded_by=self.request.user)
            .select_related('paper')
            .order_by('-created_at')
        )


@login_required
def summary_detail_redirect(request, pk):
    """
    Legacy URL: /summarizer/<summary_pk>/ → 301 to /papers/<paper_pk>/.

    The summary detail page was consolidated into the paper detail page;
    this view keeps existing bookmarks and external links working.
    Returns 404 if the summary doesn't belong to the requesting user.
    """
    from django.shortcuts import redirect
    summary = get_object_or_404(
        Summary, pk=pk, paper__uploaded_by=request.user,
    )
    return redirect('papers:detail', pk=summary.paper_id, permanent=True)


# ---- Share Links ---------------------------------------------------------

@login_required
def create_share_link(request, pk):
    """Create a new share link for a summary."""
    summary = get_object_or_404(Summary, pk=pk, paper__uploaded_by=request.user)

    if request.method == 'POST':
        expiry_days = int(request.POST.get('expiry_days', 7))
        expires_at = None
        if expiry_days > 0:
            expires_at = timezone.now() + timezone.timedelta(days=expiry_days)

        link = ShareLink.objects.create(summary=summary, expires_at=expires_at)
        messages.success(request, _('Share link created successfully.'))
        return redirect('papers:detail', pk=summary.paper_id)

    return redirect('papers:detail', pk=summary.paper_id)


@login_required
def revoke_share_link(request, link_pk):
    """Deactivate a share link."""
    link = get_object_or_404(ShareLink, pk=link_pk, summary__paper__uploaded_by=request.user)
    link.is_active = False
    link.save(update_fields=['is_active'])
    messages.success(request, _('Share link revoked.'))
    return redirect('papers:detail', pk=link.summary.paper_id)


def shared_summary_view(request, token):
    """Public view of a shared summary — no login required."""
    link = get_object_or_404(ShareLink, token=token)

    if not link.is_valid():
        return render(request, 'summarizer/share_expired.html', status=410)

    link.access_count += 1
    link.save(update_fields=['access_count'])

    summary = link.summary
    return render(request, 'summarizer/shared_summary.html', {
        'summary':    summary,
        'key_points': summary.get_key_points_list(),
        'citations':  summary.get_citations_list(),
        'link':       link,
    })


# ---- Chat with paper -----------------------------------------------------

@login_required
def chat_history(request, pk):
    """Return the user's chat history for this paper as JSON."""
    paper = get_object_or_404(Paper, pk=pk, uploaded_by=request.user)
    msgs = ChatMessage.objects.filter(paper=paper, user=request.user).order_by('created_at')
    return JsonResponse({
        'messages': [
            {
                'role':       m.role,
                'content':    m.content,
                'created_at': m.created_at.isoformat(),
            }
            for m in msgs
        ],
    })


@login_required
@require_POST
def chat_send(request, pk):
    """Accept a new user message, call Groq, persist + return the assistant reply."""
    paper = get_object_or_404(Paper, pk=pk, uploaded_by=request.user)

    if not hasattr(paper, 'summary'):
        return JsonResponse(
            {'error': 'This paper has not been summarized yet — generate a summary first.'},
            status=400,
        )

    try:
        payload = json.loads(request.body.decode('utf-8'))
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse({'error': 'Invalid JSON body.'}, status=400)

    content = (payload.get('content') or '').strip()
    if not content:
        return JsonResponse({'error': 'Message is empty.'}, status=400)
    if len(content) > 2000:
        return JsonResponse({'error': 'Message too long (max 2000 chars).'}, status=400)

    # Persist user message first so it survives even if the API call fails
    user_msg = ChatMessage.objects.create(
        paper=paper, user=request.user, role='user', content=content,
    )

    history = list(
        ChatMessage.objects
        .filter(paper=paper, user=request.user)
        .exclude(pk=user_msg.pk)
        .order_by('created_at')
    )

    try:
        from .services import chat_with_paper
        reply = chat_with_paper(paper, paper.summary, history, content)
    except ValueError as exc:
        logger.error('chat_with_paper failed for paper %s: %s', pk, exc)
        return JsonResponse({'error': str(exc)}, status=500)
    except Exception as exc:
        logger.exception('Unexpected chat error for paper %s', pk)
        return JsonResponse({'error': f'Unexpected error: {exc}'}, status=500)

    assistant_msg = ChatMessage.objects.create(
        paper=paper, user=request.user, role='assistant', content=reply,
    )

    return JsonResponse({
        'reply':      reply,
        'message_id': assistant_msg.pk,
        'created_at': assistant_msg.created_at.isoformat(),
    })


@login_required
@require_POST
def chat_clear(request, pk):
    """Wipe the chat history for this paper for this user."""
    paper = get_object_or_404(Paper, pk=pk, uploaded_by=request.user)
    deleted, _ = ChatMessage.objects.filter(paper=paper, user=request.user).delete()
    return JsonResponse({'deleted': deleted})


# ---- PDF / Word exports --------------------------------------------------

@login_required
def export_summary_pdf(request, pk):
    summary = get_object_or_404(Summary, pk=pk, paper__uploaded_by=request.user)

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=A4,
            leftMargin=2.5*cm, rightMargin=2.5*cm,
            topMargin=2.5*cm, bottomMargin=2.5*cm,
        )
        styles = getSampleStyleSheet()
        title_style   = ParagraphStyle('T', parent=styles['Title'], fontSize=18, spaceAfter=12,
                                       textColor=colors.HexColor('#1a1a2e'))
        heading_style = ParagraphStyle('H', parent=styles['Heading2'], fontSize=13, spaceAfter=6,
                                       spaceBefore=14, textColor=colors.HexColor('#16213e'))
        body_style    = ParagraphStyle('B', parent=styles['Normal'], fontSize=10, leading=15, spaceAfter=6)

        story = [
            Paragraph(summary.paper.title, title_style),
            Paragraph(f'Author(s): {summary.paper.author or "N/A"}', body_style),
            Spacer(1, 0.3*cm),
        ]

        for heading, content in [
            ('Abstract',    summary.abstract),
            ('Methodology', summary.methodology),
            ('Results',     summary.results),
            ('Conclusion',  summary.conclusion),
        ]:
            if content:
                story += [Paragraph(heading, heading_style), Paragraph(content.replace('\n', '<br/>'), body_style)]

        kps = summary.get_key_points_list()
        if kps:
            story.append(Paragraph('Key Points', heading_style))
            story.append(ListFlowable(
                [ListItem(Paragraph(pt, body_style), bulletColor=colors.HexColor('#4361ee')) for pt in kps],
                bulletType='bullet',
            ))

        citations = summary.get_citations_list()
        if citations:
            story.append(Paragraph('References', heading_style))
            for i, c in enumerate(citations, 1):
                story.append(Paragraph(f'[{i}] {c}', body_style))

        doc.build(story)
        buffer.seek(0)

        safe = ''.join(c for c in summary.paper.title[:50] if c.isalnum() or c in ' -_').strip()
        response = HttpResponse(buffer.read(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="summary_{safe or summary.pk}.pdf"'
        return response

    except Exception as exc:
        logger.error('PDF export failed for summary %s: %s', pk, exc)
        return HttpResponse(f'Export failed: {exc}', status=500)


@login_required
def export_summary_word(request, pk):
    summary = get_object_or_404(Summary, pk=pk, paper__uploaded_by=request.user)

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        t = doc.add_heading(summary.paper.title, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if summary.paper.author:
            a = doc.add_paragraph(f'Author(s): {summary.paper.author}')
            a.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        for heading, content in [
            ('Abstract',    summary.abstract),
            ('Methodology', summary.methodology),
            ('Results',     summary.results),
            ('Conclusion',  summary.conclusion),
        ]:
            if content:
                doc.add_heading(heading, level=1)
                doc.add_paragraph(content)

        kps = summary.get_key_points_list()
        if kps:
            doc.add_heading('Key Points', level=1)
            for pt in kps:
                doc.add_paragraph(pt, style='List Bullet')

        citations = summary.get_citations_list()
        if citations:
            doc.add_heading('References', level=1)
            for i, c in enumerate(citations, 1):
                doc.add_paragraph(f'[{i}] {c}')

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        safe = ''.join(c for c in summary.paper.title[:50] if c.isalnum() or c in ' -_').strip()
        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        response['Content-Disposition'] = f'attachment; filename="summary_{safe or summary.pk}.docx"'
        return response

    except Exception as exc:
        logger.error('Word export failed for summary %s: %s', pk, exc)
        return HttpResponse(f'Export failed: {exc}', status=500)
